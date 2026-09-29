"""Acceptance tests for firm.paydata on a synthetic LCA workbook written at test time."""
import pathlib
import sys

import numpy as np
import openpyxl

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1]))
import firm_paydata as pd  # noqa: E402

HEADER = ["CASE_NUMBER", "CASE_STATUS", "VISA_CLASS", "JOB_TITLE", "SOC_CODE", "EMPLOYER_NAME", "EMPLOYER_POC_EMAIL",
          "WAGE_RATE_OF_PAY_FROM", "WAGE_UNIT_OF_PAY", "PW_WAGE_LEVEL", "FULL_TIME_POSITION"]


def workbook(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADER)
    ws.append(["I-1", "Certified", "H-1B", "Quantitative Researcher", "15-2041.00", "EXAMPLE TRADING LLC", "x@y",
               200000, "Year", "II", "Y"])
    ws.append(["I-2", "Denied", "H-1B", "Software Engineer", "15-1252.00", "EXAMPLE TRADING LLC", "x@y", 90, "Hour",
               "I", "Y"])
    ws.append(["I-3", "Certified", "E-3 Australian", "Associate", "15-1252.00", "OTHER BANK N.A.", "x@y", "10,000",
               "Month", "IV", "Y"])
    p = tmp_path / "lca.xlsx"
    wb.save(p)
    return p


def test_stream_and_filter(tmp_path):
    rows = list(pd.read_lca(workbook(tmp_path), keep=lambda d: d["CASE_STATUS"] == "Certified"))
    assert [r["CASE_NUMBER"] for r in rows] == ["I-1", "I-3"]
    assert "EMPLOYER_POC_EMAIL" not in rows[0]


def test_annualise():
    assert pd.annualise(90, "Hour") == 90 * 2080 and pd.annualise("10,000", "Month") == 120_000
    assert pd.annualise(200000, "Year") == 200_000 and pd.annualise("", "Year") is None
    assert pd.annualise(5, "Fortnight") is None


def test_classifiers(tmp_path):
    rules = [pd.Rule("^EXAMPLE TRADING", "Example", "market maker", "x")]
    assert pd.classify("Example Trading LLC", rules).kind == "market maker"
    assert pd.classify("Other Bank N.A.", rules) is None
    assert pd.classify_role("Quantitative Researcher", "") == "quant researcher"
    assert pd.classify_role("Senior Quant Trader", "") == "trader"
    assert pd.classify_role("Associate", "15-1252.00") == "software engineer"
    assert pd.classify_role("Portfolio Manager, Quantitative Research", "") == "portfolio manager"
    assert pd.classify_role("Chef", "35-1011.00") is None
    p = tmp_path / "r.csv"
    p.write_text("pattern,label,kind,ledger\n^A ,A,bank,F1\n")
    assert pd.load_rules(p)[0].label == "A"


def test_cells_and_suppression():
    rng = np.random.default_rng(0)
    assert pd.cell(range(9)) == {"n": 9, "suppressed": True}
    c = pd.cell([float(x) for x in range(1, 101)], rng, 500)
    assert c["n"] == 100 and abs(c["p50"] - 50.5) < 1e-9 and c["lo50"] < 50.5 < c["hi50"]
    assert c["p10"] < c["p25"] < c["p50"] < c["p75"] < c["p90"]


def test_evidence_record():
    e = pd.PayEvidence("s", "m", "p", 2025, 1.0, 2.0, "USD", "c")
    assert e.lo <= e.hi and e.caveat == "c"
